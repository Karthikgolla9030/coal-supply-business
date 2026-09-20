import io
import csv
from decimal import Decimal
from django.db.models import Sum, Q, F, Value, CharField, DateField
from django.db.models.functions import Coalesce, Cast
from django.utils.timezone import now
from django.template.loader import render_to_string
from django.conf import settings
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import xhtml2pdf.pisa as pisa

from core.models import LedgerEntry, LedgerPayment, Invoice, Customer


def get_report_data(business, report_type, start_date=None, end_date=None, party_id=None, payment_method=None, include_voided=False):
    """
    Fetches structured data for the requested report type.
    Returns (summary_dict, rows_list)
    """
    summary = {
        "to_receive": Decimal("0.00"),
        "to_pay": Decimal("0.00"),
        "received": Decimal("0.00"),
        "paid": Decimal("0.00"),
        "current_outstanding_receive": Decimal("0.00"),
        "current_outstanding_pay": Decimal("0.00")
    }

    # Calculate Current Outstanding (Independent of date range as per requirements)
    all_active_entries = LedgerEntry.objects.filter(business=business).exclude(status='VOIDED')
    for e in all_active_entries:
        rem = e.get_remaining_amount()
        if rem > 0:
            if e.transaction_type == "RECEIVABLE":
                summary["current_outstanding_receive"] += rem
            else:
                summary["current_outstanding_pay"] += rem

    # Date filters for activity
    entry_filter = Q(business=business)
    payment_filter = Q(ledger_entry__business=business)
    
    if not include_voided:
        entry_filter &= ~Q(status='VOIDED')
        payment_filter &= Q(status='RECORDED')
    
    if start_date:
        entry_filter &= Q(created_at__date__gte=start_date)
        payment_filter &= Q(payment_date__gte=start_date)
    if end_date:
        entry_filter &= Q(created_at__date__lte=end_date)
        payment_filter &= Q(payment_date__lte=end_date)

    if party_id:
        entry_filter &= Q(customer_id=party_id)
        payment_filter &= Q(ledger_entry__customer_id=party_id)
        
    if payment_method:
        payment_filter &= Q(payment_method=payment_method)

    rows = []

    if report_type == "Financial Summary":
        # Calculate received and paid during the selected period
        aggr_payments = LedgerPayment.objects.filter(payment_filter).annotate(
            ttype=F("ledger_entry__transaction_type")
        ).values("ttype").annotate(total=Sum("amount"))
        
        for ap in aggr_payments:
            if ap["ttype"] == "RECEIVABLE":
                summary["received"] = ap["total"] or Decimal("0.00")
            elif ap["ttype"] == "PAYABLE":
                summary["paid"] = ap["total"] or Decimal("0.00")
                
        # Calculate money to receive/pay generated during the selected period
        aggr_entries = LedgerEntry.objects.filter(entry_filter).values("transaction_type").annotate(total=Sum("amount"))
        for ae in aggr_entries:
            if ae["transaction_type"] == "RECEIVABLE":
                summary["to_receive"] = ae["total"] or Decimal("0.00")
            elif ae["transaction_type"] == "PAYABLE":
                summary["to_pay"] = ae["total"] or Decimal("0.00")
                
        # Financial summary has no rows, just summary
        rows = []
        
    elif report_type == "Money to Receive":
        entries = LedgerEntry.objects.filter(entry_filter, transaction_type="RECEIVABLE").order_by("-created_at")
        for e in entries:
            orig = e.amount
            paid = e.get_paid_amount()
            rem = e.get_remaining_amount()
            summary["to_receive"] += orig
            summary["received"] += paid
            rows.append({
                "Date": e.created_at.date().isoformat(),
                "Party": e.customer.name if e.customer else e.party_name,
                "Invoice / Reference": e.reference or "-",
                "Original Amount": orig,
                "Paid": paid,
                "Remaining": rem,
                "Status": e.status
            })

    elif report_type == "Money to Pay":
        entries = LedgerEntry.objects.filter(entry_filter, transaction_type="PAYABLE").order_by("-created_at")
        for e in entries:
            orig = e.amount
            paid = e.get_paid_amount()
            rem = e.get_remaining_amount()
            summary["to_pay"] += orig
            summary["paid"] += paid
            rows.append({
                "Date": e.created_at.date().isoformat(),
                "Party": e.customer.name if e.customer else e.party_name,
                "Reference": e.reference or "-",
                "Original Amount": orig,
                "Paid": paid,
                "Remaining": rem,
                "Status": e.status
            })

    elif report_type == "Payment History":
        payments = LedgerPayment.objects.filter(payment_filter).order_by("-payment_date", "-created_at")
        for p in payments:
            ttype = p.ledger_entry.transaction_type
            if ttype == "RECEIVABLE":
                summary["received"] += p.amount
            else:
                summary["paid"] += p.amount
                
            rows.append({
                "Date": p.payment_date.isoformat(),
                "Party": p.ledger_entry.customer.name if p.ledger_entry.customer else p.ledger_entry.party_name,
                "Payment Type": "Money Received" if ttype == "RECEIVABLE" else "Money Paid",
                "Invoice / Reference": p.ledger_entry.reference or "-",
                "Payment Method": p.payment_method or "Other",
                "Amount": p.amount
            })

    elif report_type == "Customer Ledger":
        entries = LedgerEntry.objects.filter(entry_filter, transaction_type="RECEIVABLE").order_by("created_at")
        payments = LedgerPayment.objects.filter(payment_filter, ledger_entry__transaction_type="RECEIVABLE").order_by("payment_date", "created_at")
        
        timeline = []
        for e in entries:
            timeline.append({
                "Date": e.created_at.date().isoformat(),
                "Description": "Invoice / Entry",
                "Reference": e.reference or "-",
                "Amount": e.amount,
                "Paid": Decimal("0.00"),
                "Balance": Decimal("0.00"),
                "_raw_date": e.created_at.date(),
                "_type": "INVOICE"
            })
            summary["to_receive"] += e.amount
        
        for p in payments:
            timeline.append({
                "Date": p.payment_date.isoformat(),
                "Description": f"Payment ({p.payment_method})",
                "Reference": p.ledger_entry.reference or "-",
                "Amount": Decimal("0.00"),
                "Paid": p.amount,
                "Balance": Decimal("0.00"),
                "_raw_date": p.payment_date,
                "_type": "PAYMENT"
            })
            summary["received"] += p.amount
            
        timeline.sort(key=lambda x: (x["_raw_date"], 1 if x["_type"] == "PAYMENT" else 0))
        
        running_balance = Decimal("0.00")
        for t in timeline:
            running_balance += t["Amount"]
            running_balance -= t["Paid"]
            t["Balance"] = running_balance
            rows.append({k: v for k, v in t.items() if not k.startswith("_")})
            
    elif report_type == "Transaction History":
        entries = LedgerEntry.objects.filter(entry_filter)
        payments = LedgerPayment.objects.filter(payment_filter)
        
        timeline = []
        for e in entries:
            timeline.append({
                "Date": e.created_at.date().isoformat(),
                "Party": e.customer.name if e.customer else e.party_name,
                "Type": "Receive" if e.transaction_type == "RECEIVABLE" else "Pay",
                "Reference": e.reference or "-",
                "Amount": e.amount,
                "Status": e.status,
                "_raw_date": e.created_at.date(),
                "_id": f"E{e.id}"
            })
            if e.transaction_type == "RECEIVABLE":
                summary["to_receive"] += e.amount
            else:
                summary["to_pay"] += e.amount
                
        for p in payments:
            ttype = p.ledger_entry.transaction_type
            timeline.append({
                "Date": p.payment_date.isoformat(),
                "Party": p.ledger_entry.customer.name if p.ledger_entry.customer else p.ledger_entry.party_name,
                "Type": "Payment Received" if ttype == "RECEIVABLE" else "Payment Made",
                "Reference": p.ledger_entry.reference or "-",
                "Amount": p.amount,
                "Status": "RECORDED",
                "_raw_date": p.payment_date,
                "_id": f"P{p.id}"
            })
            if ttype == "RECEIVABLE":
                summary["received"] += p.amount
            else:
                summary["paid"] += p.amount
                
        timeline.sort(key=lambda x: (x["_raw_date"], x["_id"]), reverse=True)
        for t in timeline:
            rows.append({k: v for k, v in t.items() if not k.startswith("_")})

    return summary, rows


def generate_pdf_report(business, report_type, date_range_label, summary, rows):
    context = {
        "business": business,
        "report_type": report_type,
        "date_range_label": date_range_label,
        "summary": summary,
        "rows": rows,
        "generated_date": now().date().isoformat()
    }
    html_string = render_to_string("report_pdf.html", context)
    pdf_file = io.BytesIO()
    pisa_status = pisa.CreatePDF(html_string, dest=pdf_file)
    if pisa_status.err:
        raise Exception("Error generating PDF report")
    return pdf_file.getvalue()


def generate_excel_report(business, report_type, date_range_label, summary, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = report_type[:31]  # Excel sheet names max 31 chars
    
    ws.append([business.business_name.upper()])
    ws.append(["FINANCIAL REPORT"])
    ws.append([f"Report Type: {report_type}"])
    ws.append([f"Period: {date_range_label}"])
    ws.append([f"Generated: {now().date().isoformat()}"])
    ws.append([])
    
    header_font = Font(bold=True)
    for i in range(1, 6):
        ws.cell(row=i, column=1).font = header_font
        
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        
        row_num = 7
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="EEEEEE", end_color="EEEEEE", fill_type="solid")
            
        for r_idx, row_dict in enumerate(rows, start=8):
            for c_idx, key in enumerate(headers, start=1):
                val = row_dict[key]
                if isinstance(val, Decimal):
                    cell = ws.cell(row=r_idx, column=c_idx, value=float(val))
                    cell.number_format = '"₹"#,##0.00'
                else:
                    ws.cell(row=r_idx, column=c_idx, value=str(val))
                    
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
    else:
        ws.append(["No records found for the selected period."])
        
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def generate_csv_report(rows):
    out = io.StringIO()
    if rows:
        headers = list(rows[0].keys())
        writer = csv.DictWriter(out, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            formatted_row = {}
            for k, v in row.items():
                if isinstance(v, Decimal):
                    formatted_row[k] = f"{v:.2f}"
                else:
                    formatted_row[k] = v
            writer.writerow(formatted_row)
    else:
        out.write("No records found.\n")
    return out.getvalue()
